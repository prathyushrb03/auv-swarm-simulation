#!/usr/bin/env python

from openpyxl import load_workbook
from numpy import interp

# target PWM
pwm = [1100, 1150, 1200, 1250, 1300, 1350, 1400, 1450, 1480, 1520, 1550, 1600, 1650, 1700, 1750, 1800, 1850, 1900]

wb = load_workbook('T200-Public-Performance-Data-10-20V-September-2019.xlsx',data_only=True)
ws = wb['16 V']

def cell(row, col):
    return ws.cell(row=row, column=col).value

rows = range(4, 205)
pwm_col = 1
vel_col = 8
thrust_col = 11

pwms = [cell(row, pwm_col) for row in rows]
vels = [cell(row, vel_col) for row in rows]
thrusts = [cell(row, thrust_col) for row in rows]

def cpp(arr, name):
    print(f'constexpr auto {name}{{std::array{{',
        ','.join(f'{v:.02f}f' for v in arr),
    '}};')

cpp(pwm, 'pwm')

for name,dst in (('vels',vels), ('thrusts', thrusts)):
    cpp([interp(p, pwms, dst) for p in pwm], name)
